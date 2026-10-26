"""Text extraction for PDF, DOCX and TXT. Each parser returns a list of pages
(DOCX/TXT are a single page) so the cleaner can strip repeated headers/footers."""
from pathlib import Path

import fitz  # PyMuPDF
import pdfplumber
from docx import Document as DocxDocument

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MIN_PDF_CHARS = 50  # below this we assume the first parser failed (or the PDF is scanned)


class ParseError(Exception):
    pass


def validate_signature(ext: str, head: bytes) -> None:
    """Check magic bytes so a renamed file is rejected early."""
    if ext == ".pdf" and not head.lstrip().startswith(b"%PDF"):
        raise ParseError("File is not a valid PDF")
    if ext == ".docx" and not head.startswith(b"PK"):
        raise ParseError("File is not a valid DOCX")
    if ext == ".txt" and b"\x00" in head:
        raise ParseError("File does not look like plain text")


def _parse_pdf_pymupdf(path: Path) -> list[str]:
    with fitz.open(path) as doc:
        return [page.get_text("text") for page in doc]


def _parse_pdf_pdfplumber(path: Path) -> list[str]:
    with pdfplumber.open(path) as pdf:
        return [page.extract_text() or "" for page in pdf.pages]


def parse_pdf(path: Path) -> list[str]:
    pages: list[str] = []
    try:
        pages = _parse_pdf_pymupdf(path)
    except Exception:
        pages = []
    if sum(len(p.strip()) for p in pages) < MIN_PDF_CHARS:
        try:
            fallback = _parse_pdf_pdfplumber(path)
        except Exception as exc:
            if not pages:
                raise ParseError(f"Could not read PDF: {exc}") from exc
            fallback = []
        if sum(len(p.strip()) for p in fallback) > sum(len(p.strip()) for p in pages):
            pages = fallback
    if sum(len(p.strip()) for p in pages) < MIN_PDF_CHARS:
        raise ParseError("No extractable text found (scanned PDF? OCR is not supported yet)")
    return pages


def parse_docx(path: Path) -> list[str]:
    try:
        doc = DocxDocument(str(path))
    except Exception as exc:
        raise ParseError(f"Could not read DOCX: {exc}") from exc
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return ["\n".join(lines)]


def parse_txt(path: Path) -> list[str]:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return [data.decode(encoding)]
        except UnicodeError:
            continue
    raise ParseError("Could not decode text file")


def extract_pages(path: Path) -> list[str]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return parse_pdf(path)
    if ext == ".docx":
        return parse_docx(path)
    if ext == ".txt":
        return parse_txt(path)
    raise ParseError(f"Unsupported file type: {ext}")
